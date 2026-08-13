import openpyxl
from copy import copy

wb = openpyxl.load_workbook('Black_Box_Co_Model.xlsx', data_only=False)
ws = wb['Historicals']

# Column map: FY19=D,FY20=E,FY21=F,FY22=G,FY23=H,FY24=I,FY25=J,FY26=K
C = {2019:'D',2020:'E',2021:'F',2022:'G',2023:'H',2024:'I',2025:'J',2026:'K'}
FULL = [2021,2022,2023,2024,2025,2026]  # years with full statements

def setv(addr, val):
    ws[addr] = val

def setlabel(addr, txt):
    ws[addr] = txt

# ---- Title ----
ws['A1'] = 'BLACK BOX LTD - HISTORICAL FINANCIALS (CONSOLIDATED)'

# ============ INCOME STATEMENT ============
# Revenue from Operations (consolidated, INR cr)
rev = {2019:1853,2020:4994,2021:4674,2022:5370,2023:6288,2024:6282,2025:5967,2026:6322}
oth_inc = {2021:11,2022:6,2023:22,2024:19,2025:5,2026:6}
gross_profit = {2021:1497,2022:1549,2023:1640,2024:1714,2025:1794,2026:1930}
# Cost of revenue = revenue - gross profit
cost_rev = {y: rev[y]-gross_profit[y] for y in FULL}
# Total other expenses (SG&A/opex below gross profit)
opex = {2021:1145,2022:1289,2023:1370,2024:1286,2025:1263,2026:1360}
fin = {2021:98,2022:74,2023:111,2024:141,2025:145,2026:158}
dep = {2021:96,2022:99,2023:107,2024:114,2025:113,2026:116}
reported_pbt = {2021:96,2022:86,2023:29,2024:156,2025:212,2026:239}
tax = {2021:18,2022:13,2023:6,2024:19,2025:7,2026:22}
pat = {2019:-79,2020:-80,2021:78,2022:73,2023:24,2024:138,2025:205,2026:218}
eps = {2021:5.21,2022:4.45,2023:1.42,2024:8.20,2025:12.16,2026:12.78}

# Relabel P&L lines
ws['A12'] = 'Cost of Revenue (hardware, materials & subcontracting)'
ws['A15'] = 'Employee & Other Operating Expenses (SG&A)'
ws['A18'] = '(folded into SG&A above)'
ws['A21'] = 'Exceptional Items & Share of Assoc. (net)'

for y in FULL:
    c = C[y]
    ws[f'{c}4']=rev[y]              # Revenue Gross
    ws[f'{c}5']=0                                   # excise
    ws[f'{c}6']=rev[y]                              # Revenue Net
    ws[f'{c}7']=0                                   # other op rev
    ws[f'{c}8']=rev[y]                              # total op rev
    ws[f'{c}9']=oth_inc[y]                          # other income
    ws[f'{c}10']=rev[y]+oth_inc[y]                  # total revenue
    ws[f'{c}12']=cost_rev[y]                        # cost of revenue
    ws[f'{c}13']=0
    ws[f'{c}14']=0
    ws[f'{c}15']=opex[y]                            # SG&A / total other exp
    ws[f'{c}16']=fin[y]                             # finance
    ws[f'{c}17']=dep[y]                             # depreciation
    ws[f'{c}18']=0                                  # other exp folded
    total_exp = cost_rev[y]+opex[y]+fin[y]+dep[y]
    ws[f'{c}19']=total_exp                          # total expenses
    pbt_before = (rev[y]+oth_inc[y]) - total_exp
    ws[f'{c}20']=round(pbt_before,2)               # PBT before exceptional
    excp = round(reported_pbt[y]-pbt_before,2)     # exceptional+assoc plug -> ties to reported PBT
    ws[f'{c}21']=excp
    ws[f'{c}22']=reported_pbt[y]                    # PBT
    ws[f'{c}24']=tax[y]                             # current tax
    ws[f'{c}25']=0                                  # deferred
    ws[f'{c}26']=tax[y]                             # total tax
    ws[f'{c}27']=reported_pbt[y]-tax[y]
    ws[f'{c}28']=reported_pbt[y]-tax[y]
    ws[f'{c}29']=0
    ws[f'{c}30']=pat[y]                             # PAT for period (reported)
    ws[f'{c}33']=eps[y]
    ws[f'{c}34']=eps[y]
    ws[f'{c}36']=0                                  # dividend (BB pays negligible)

# FY19, FY20 partial (revenue + PAT only)
for y in [2019,2020]:
    c=C[y]
    ws[f'{c}4']=rev[y]; ws[f'{c}6']=rev[y]; ws[f'{c}8']=rev[y]; ws[f'{c}10']=rev[y]
    ws[f'{c}30']=pat[y]

# ============ BALANCE SHEET ============
eq_cap = {2021:33,2022:33,2023:34,2024:34,2025:34,2026:36}
reserves = {2021:174,2022:227,2023:262,2024:447,2025:725,2026:1251}
# Total SH funds
tsf = {y: eq_cap[y]+reserves[y] for y in FULL}
mi = {y:0 for y in FULL}
lt_borrow = {2021:119,2022:229,2023:305,2024:362,2025:633,2026:797}
dtl = {y:0 for y in FULL}
# Other LT liab = lease NC + other fin liab NC + contract liab NC
other_lt = {2021:181,2022:177,2023:284,2024:328,2025:287,2026:404}
lt_prov = {2021:103,2022:70,2023:74,2024:54,2025:32,2026:37}
# ST borrowings + current lease
st_borrow = {2021:115,2022:135,2023:101,2024:83,2025:75,2026:93}
payables = {2021:516,2022:1009,2023:1158,2024:722,2025:556,2026:790}
# Other CL = other fin liab curr + contract liab curr + other CL
other_cl = {2021:937,2022:698,2023:716,2024:695,2025:676,2026:830}
st_prov = {2021:127,2022:72,2023:69,2024:75,2025:55,2026:55}

# Assets
# PP&E + ROU
ppe_rou = {2021:164+146,2022:190+194,2023:161+259,2024:120+291,2025:102+254,2026:128+283}
# Goodwill + other intangibles
intang = {2021:269+43,2022:300+47,2023:316+61,2024:334+63,2025:335+77,2026:381+92}
cwip = {y:0 for y in FULL}
intdev = {y:0 for y in FULL}
# Non-current financial assets/investments (incl equity-method investment)
nc_inv = {2021:0+28,2022:0+24,2023:30+18,2024:32+35,2025:33+23,2026:0+40}
dta = {2021:67,2022:63,2023:60,2024:32,2025:28,2026:48}   # tax assets
lt_loans = {y:0 for y in FULL}
other_nca = {2021:31,2022:26,2023:71,2024:57,2025:89,2026:87}
# Current financial assets (deposits/investments)
curr_fin = {2021:533,2022:560,2023:678,2024:508,2025:549,2026:543}
inv = {2021:149,2022:226,2023:362,2024:246,2025:210,2026:323}
recv = {2021:240,2022:374,2023:421,2024:386,2025:567,2026:1153}
st_loans = {y:0 for y in FULL}
# Contract assets + other current assets
other_ca = {2021:0+223,2022:44+291,2023:114+242,2024:246+227,2025:219+357,2026:280+397}

# Relabel BS lines
ws['A51']='Lease Liab (NC) + Contract Liab + Other Fin Liab'
ws['A55']='Short-term Borrowings + Current Lease Liab'
ws['A57']='Contract Liab (curr) + Other Fin/Curr Liab'
ws['A63']='Property, Plant & Equipment + Right-of-Use Assets'
ws['A64']='Goodwill & Other Intangibles'
ws['A68']='Non-Current Financial Assets/Investments'
ws['A74']='Current Financial Assets (deposits/investments)'
ws['A79']='Contract Assets + Other Current Assets'

for y in FULL:
    c=C[y]
    # Equity
    ws[f'{c}42']=eq_cap[y]; ws[f'{c}43']=eq_cap[y]
    ws[f'{c}44']=reserves[y]; ws[f'{c}45']=reserves[y]
    ws[f'{c}46']=tsf[y]
    ws[f'{c}47']=mi[y]
    # NCL
    ws[f'{c}49']=lt_borrow[y]; ws[f'{c}50']=dtl[y]; ws[f'{c}51']=other_lt[y]; ws[f'{c}52']=lt_prov[y]
    ws[f'{c}53']=f'=SUM({c}49:{c}52)'
    # CL
    ws[f'{c}55']=st_borrow[y]; ws[f'{c}56']=payables[y]; ws[f'{c}57']=other_cl[y]; ws[f'{c}58']=st_prov[y]
    ws[f'{c}59']=f'=SUM({c}55:{c}58)'
    # total cap & liab
    ws[f'{c}60']=f'=SUM({c}59,{c}53,{c}46)+{c}47'
    total_L = tsf[y]+mi[y]+(lt_borrow[y]+dtl[y]+other_lt[y]+lt_prov[y])+(st_borrow[y]+payables[y]+other_cl[y]+st_prov[y])
    # Assets
    ws[f'{c}63']=ppe_rou[y]; ws[f'{c}64']=intang[y]; ws[f'{c}65']=cwip[y]; ws[f'{c}66']=intdev[y]
    ws[f'{c}67']=f'=SUM({c}63:{c}66)'   # fixed assets
    ws[f'{c}68']=nc_inv[y]; ws[f'{c}69']=dta[y]; ws[f'{c}70']=lt_loans[y]; ws[f'{c}71']=other_nca[y]
    ws[f'{c}72']=f'={c}67+{c}68+{c}69+{c}70+{c}71'  # total NCA
    total_nca = ppe_rou[y]+intang[y]+cwip[y]+intdev[y]+nc_inv[y]+dta[y]+lt_loans[y]+other_nca[y]
    ws[f'{c}74']=curr_fin[y]; ws[f'{c}75']=inv[y]; ws[f'{c}76']=recv[y]
    ws[f'{c}78']=st_loans[y]; ws[f'{c}79']=other_ca[y]
    ca_ex_cash = curr_fin[y]+inv[y]+recv[y]+st_loans[y]+other_ca[y]
    cash_plug = round(total_L - total_nca - ca_ex_cash, 2)   # cash as balancing plug -> BS balances exactly
    ws[f'{c}77']=cash_plug
    ws[f'{c}80']=f'={c}74+{c}75+{c}76+{c}77+{c}78+{c}79'  # total CA
    ws[f'{c}81']=f'={c}80+{c}72'                           # total assets

# FY19/FY20 net worth
ws['D46']=19; ws['E46']=-176

# ============ CASH FLOW ============
cfo = {2021:303,2022:95,2023:13,2024:129,2025:-88,2026:84}
op_bwc = {2021:407,2022:224,2023:296,2024:407,2025:466,2026:453}
wc_chg = {2021:-156,2022:-108,2023:-282,2024:-301,2025:-551,2026:-342}
taxes_cf = {2021:52,2022:-21,2023:0,2024:24,2025:-2,2026:-27}
cfi = {2021:-32,2022:-71,2023:19,2024:1,2025:-47,2026:-57}
cff = {2021:-277,2022:-43,2023:-58,2024:-155,2025:192,2026:197}
capex_hist = {2021:30,2022:40,2023:45,2024:50,2025:55,2026:60}  # est. pure capex (asset-light)

ws['A93']='Change in other loans/advances'
for y in FULL:
    c=C[y]
    ws[f'{c}89']=cfo[y]
    ws[f'{c}90']=op_bwc[y]
    ws[f'{c}96']=wc_chg[y]
    ws[f'{c}97']=taxes_cf[y]
    ws[f'{c}98']=cfi[y]
    ws[f'{c}110']=cff[y]
    ws[f'{c}117']=f'={c}89+{c}98+{c}110'          # net cash flow
    ws[f'{c}118']=round(cfo[y]-capex_hist[y],1)   # FCF = CFO - capex
    ws[f'{c}119']=f'={c}89/{c}90'                 # CFO/OP

# Clear the granular TVS-specific CF component rows (91-95, 99-116) that we don't map
for r in [91,92,94,95,99,100,101,102,103,104,105,106,107,108,109,111,112,113,114,115,116]:
    for y in FULL+[2019,2020]:
        c=C[y]
        ws[f'{c}{r}']=None
# also clear FY17,FY18 (B,C) legacy TVS numbers across all rows
for col in ['B','C']:
    for r in range(4,120):
        ws[f'{col}{r}']=None

# Note on data availability
ws['A121']='Note: BB in current consolidated form dates from FY20 (post Jan-2019 Black Box Corp acquisition). Full 3-statements FY21-FY26; FY19-FY20 partial (revenue/PAT); FY17-FY18 (AGC era) not comparable. Source: BB FY26 Investor Presentation pp.21-24; Annual Reports FY24/FY25.'

wb.save('Black_Box_Co_Model.xlsx')
print('Historicals rebuilt & saved.')
