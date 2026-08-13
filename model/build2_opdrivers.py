import openpyxl
from copy import copy
wb = openpyxl.load_workbook('Black_Box_Co_Model.xlsx', data_only=False)
ws = wb['Operating Drivers']

# capture a couple style templates
red_hdr = copy(ws['A3'].font)      # red bold section header
yr_font = copy(ws['C3'].font)      # blue year header
pct_fmt = '0.0%'
num_fmt = '#,##0'
day_fmt = '0.0'

# Clear existing content A3:R36
for r in range(3,37):
    for col in range(1,19):
        ws.cell(row=r, column=col).value = None

ws['A1']='BLACK BOX LTD - OPERATING DRIVERS & KPIs (CONSOLIDATED)'

# Year headers: FY21..FY26 in G..L ; FY27E..FY32E in M..R
# OpDrivers col -> Historicals col (offset by 1): G(FY21)->F ... L(FY26)->K
od_hist = {'G':'F','H':'G','I':'H','J':'I','K':'J','L':'K'}  # FY21..FY26
od_years = ['G','H','I','J','K','L']
fc_cols = ['M','N','O','P','Q','R']   # FY27E..FY32E
fc_src  = ['D','E','F','G','H','I']   # Forecast columns FY27E..FY32E

def yr_headers(row):
    labels = {'G':'FY 21','H':'FY 22','I':'FY 23','J':'FY 24','K':'FY 25','L':'FY 26',
              'M':'FY 27E','N':'FY 28E','O':'FY 29E','P':'FY 30E','Q':'FY 31E','R':'FY 32E'}
    for col,lab in labels.items():
        c=ws[f'{col}{row}']; c.value=lab; c.font=copy(yr_font)

def section(row, text):
    c=ws[f'A{row}']; c.value=text; c.font=copy(red_hdr)

def rowlabel(row, text, unit=None):
    ws[f'A{row}']=text
    if unit: ws[f'B{row}']=unit

# ---------------- SECTION 1: REVENUE & GROWTH ----------------
section(3,'SECTION 1 — REVENUE & GROWTH (organic vs inorganic)')
yr_headers(3)
rowlabel(4,'Total revenue','INR Cr')
for col,h in od_hist.items():
    ws[f'{col}4']=f"=Historicals!{h}6"; ws[f'{col}4'].number_format=num_fmt
for i,col in enumerate(fc_cols):
    ws[f'{col}4']=f"=Forecast!{fc_src[i]}30"; ws[f'{col}4'].number_format=num_fmt
rowlabel(5,'Revenue growth % YoY','%')
allcols=od_years+fc_cols
for i in range(1,len(allcols)):
    c=allcols[i]; p=allcols[i-1]
    ws[f'{c}5']=f"={c}4/{p}4-1"; ws[f'{c}5'].number_format=pct_fmt
rowlabel(6,'Inorganic (acquired) revenue added','INR Cr')
for i,col in enumerate(fc_cols):
    ws[f'{col}6']=f"=Forecast!{fc_src[i]}5"; ws[f'{col}6'].number_format=num_fmt
for col in od_years: ws[f'{col}6']=0
rowlabel(7,'Organic revenue growth % (ex-M&A)','%')
for i,col in enumerate(fc_cols):
    ws[f'{col}7']=f"=Forecast!{fc_src[i]}4"; ws[f'{col}7'].number_format=pct_fmt
# historical organic ~ total (no M&A FY21-26)
for i in range(1,len(od_years)):
    c=od_years[i]; p=od_years[i-1]
    ws[f'{c}7']=f"={c}4/{p}4-1"; ws[f'{c}7'].number_format=pct_fmt

# ---------------- SECTION 2: SEGMENT, RECURRING & GEO MIX ----------------
section(9,'SECTION 2 — SEGMENT, RECURRING & GEOGRAPHIC MIX (disclosed)')
rowlabel(10,'Global Solutions Integration (GSI) % rev','%'); ws['L10']=0.84; ws['L10'].number_format=pct_fmt
rowlabel(11,'Technology Product Solutions (TPS) % rev','%'); ws['L11']=0.14; ws['L11'].number_format=pct_fmt
rowlabel(12,'Other services % rev','%'); ws['L12']=0.02; ws['L12'].number_format=pct_fmt
rowlabel(13,'Recurring/managed services % rev','%'); ws['K13']=0.30; ws['L13']=0.30
ws['K13'].number_format=pct_fmt; ws['L13'].number_format=pct_fmt
rowlabel(14,'Data-centre revenue % (mix)','%'); ws['L14']=0.17; ws['M14']=0.30
ws['L14'].number_format=pct_fmt; ws['M14'].number_format=pct_fmt
rowlabel(15,'Geo mix: North America / Europe / India','%'); ws['L15']='NA 69% | EU 9% | IN 6%'
rowlabel(16,'Geo mix: APAC / LatAm / MEA','%'); ws['L16']='APAC 8% | LatAm 6% | MEA 2%'

# ---------------- SECTION 3: MARGINS ----------------
section(18,'SECTION 3 — MARGIN STRUCTURE')
rowlabel(19,'Gross margin %','%')
for col,h in od_hist.items():
    ws[f'{col}19']=f"=(Historicals!{h}6-Historicals!{h}12)/Historicals!{h}6"; ws[f'{col}19'].number_format=pct_fmt
for i,col in enumerate(fc_cols):
    ws[f'{col}19']=f"=1-Forecast!{fc_src[i]}6"; ws[f'{col}19'].number_format=pct_fmt
rowlabel(20,'SG&A (total other opex) % rev','%')
for col,h in od_hist.items():
    ws[f'{col}20']=f"=Historicals!{h}15/Historicals!{h}6"; ws[f'{col}20'].number_format=pct_fmt
for i,col in enumerate(fc_cols):
    ws[f'{col}20']=f"=Forecast!{fc_src[i]}9"; ws[f'{col}20'].number_format=pct_fmt
rowlabel(21,'EBITDA margin %','%')
for col,h in od_hist.items():
    ws[f'{col}21']=f"=(Historicals!{h}6-Historicals!{h}12-Historicals!{h}15)/Historicals!{h}6"; ws[f'{col}21'].number_format=pct_fmt
for i,col in enumerate(fc_cols):
    ws[f'{col}21']=f"=Forecast!{fc_src[i]}52"; ws[f'{col}21'].number_format=pct_fmt
rowlabel(22,'EBIT margin %','%')
for col,h in od_hist.items():
    ws[f'{col}22']=f"=(Historicals!{h}22+Historicals!{h}16-Historicals!{h}21)/Historicals!{h}6"; ws[f'{col}22'].number_format=pct_fmt

# ---------------- SECTION 4: WORKING CAPITAL ----------------
section(24,'SECTION 4 — WORKING CAPITAL EFFICIENCY')
rowlabel(25,'DSO (receivable days)','days')
for col,h in od_hist.items():
    ws[f'{col}25']=f"=Historicals!{h}76/Historicals!{h}6*365"; ws[f'{col}25'].number_format=day_fmt
for i,col in enumerate(fc_cols):
    ws[f'{col}25']=f"=Forecast!{fc_src[i]}18"; ws[f'{col}25'].number_format=day_fmt
rowlabel(26,'Inventory days','days')
for col,h in od_hist.items():
    ws[f'{col}26']=f"=Historicals!{h}75/Historicals!{h}6*365"; ws[f'{col}26'].number_format=day_fmt
for i,col in enumerate(fc_cols):
    ws[f'{col}26']=f"=Forecast!{fc_src[i]}19"; ws[f'{col}26'].number_format=day_fmt
rowlabel(27,'DPO (payable days, on revenue)','days')
for col,h in od_hist.items():
    ws[f'{col}27']=f"=Historicals!{h}56/Historicals!{h}6*365"; ws[f'{col}27'].number_format=day_fmt
for i,col in enumerate(fc_cols):
    ws[f'{col}27']=f"=Forecast!{fc_src[i]}20"; ws[f'{col}27'].number_format=day_fmt
rowlabel(28,'Contract assets % of revenue','%')
for col,h in od_hist.items():
    ws[f'{col}28']=f"=Historicals!{h}79/Historicals!{h}6"; ws[f'{col}28'].number_format=pct_fmt
rowlabel(29,'Contract liab + other CL % of revenue','%')
for col,h in od_hist.items():
    ws[f'{col}29']=f"=Historicals!{h}57/Historicals!{h}6"; ws[f'{col}29'].number_format=pct_fmt

# ---------------- SECTION 5: CAPITAL, RETURNS & TAX ----------------
section(31,'SECTION 5 — CAPITAL, RETURNS & TAX')
rowlabel(32,'Net debt / (net cash)','INR Cr')
for col,h in od_hist.items():
    ws[f'{col}32']=f"=Historicals!{h}49+Historicals!{h}55-Historicals!{h}77"; ws[f'{col}32'].number_format=num_fmt
rowlabel(33,'Effective tax rate','%')
for col,h in od_hist.items():
    ws[f'{col}33']=f"=Historicals!{h}26/Historicals!{h}22"; ws[f'{col}33'].number_format=pct_fmt
for i,col in enumerate(fc_cols):
    ws[f'{col}33']=f"=Forecast!{fc_src[i]}14"; ws[f'{col}33'].number_format=pct_fmt
rowlabel(34,'ROE (PAT / avg equity)','%')
for i in range(1,len(od_years)):
    c=od_years[i]; hc=od_hist[c]; hp=od_hist[od_years[i-1]]
    ws[f'{c}34']=f"=Historicals!{hc}30/AVERAGE(Historicals!{hp}46,Historicals!{hc}46)"; ws[f'{c}34'].number_format=pct_fmt
rowlabel(35,'Capex % of revenue','%')
for i,col in enumerate(fc_cols):
    ws[f'{col}35']=f"=Forecast!{fc_src[i]}13"; ws[f'{col}35'].number_format=pct_fmt

# ---------------- SECTION 6: ORDER BOOK KPIs (disclosed) ----------------
section(36,'SECTION 6 — ORDER BOOK ($mn): FY24 backlog 470 | FY25 504 | FY26 792 | FY26 bookings 1,003; avg tenure ~12-18 months')

wb.save('Black_Box_Co_Model.xlsx')
print('Operating Drivers rebuilt & saved.')
