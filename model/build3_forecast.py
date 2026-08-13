import openpyxl
wb = openpyxl.load_workbook('Black_Box_Co_Model.xlsx', data_only=False)
ws = wb['Forecast']
cols = ['D','E','F','G','H','I']  # FY27E..FY32E

ws['A1']='BLACK BOX LTD - 3 STATEMENT FORECAST (CONSOLIDATED)'

# ---------- A. ASSUMPTIONS (rows 4-21) ----------
ws['A4']='Organic revenue growth %'
org_g = {'D':0.13,'E':0.13,'F':0.12,'G':0.11,'H':0.10,'I':0.09}
for c,v in org_g.items(): ws[f'{c}4']=v

ws['A5']='Inorganic (acquired) revenue added (INR cr)'
inorg = {'D':500,'E':400,'F':400,'G':300,'H':300,'I':200}
for c,v in inorg.items(): ws[f'{c}5']=v; ws[f'{c}5'].number_format='#,##0'

ws['A6']='Cost of revenue % of revenue'
costp = {'D':0.695,'E':0.694,'F':0.693,'G':0.692,'H':0.692,'I':0.691}
for c,v in costp.items(): ws[f'{c}6']=v

ws['A7']='Purchases % of revenue (n/a - folded)'
for c in cols: ws[f'{c}7']=0
ws['A8']='Changes in inventories % (n/a)'
for c in cols: ws[f'{c}8']=0

ws['A9']='Total operating expenses (SG&A) % of revenue'
sga = {'D':0.212,'E':0.210,'F':0.208,'G':0.207,'H':0.205,'I':0.205}
for c,v in sga.items(): ws[f'{c}9']=v

ws['A10']='Other expenses % (n/a - folded)'
for c in cols: ws[f'{c}10']=0

ws['A11']='Cost of funds on avg borrowings (blended, incl. leases)'
ws['D11']=0.12
for c in ['E','F','G','H','I']:
    prev = cols[cols.index(c)-1]
    ws[f'{c}11']=f'={prev}11'

ws['A12']='Depreciation % of prior fixed assets'
for c in cols: ws[f'{c}12']=0.13

ws['A13']='Capex % of revenue'
capx = {'D':0.018,'E':0.017,'F':0.016,'G':0.016,'H':0.015,'I':0.015}
for c,v in capx.items(): ws[f'{c}13']=v

ws['A14']='Effective tax rate (NOL depletion -> normalising)'
taxr = {'D':0.12,'E':0.14,'F':0.16,'G':0.18,'H':0.19,'I':0.20}
for c,v in taxr.items(): ws[f'{c}14']=v

ws['A15']='Dividend payout (% PAT)'
for c in cols: ws[f'{c}15']=0.05

ws['A16']='Non-current financial assets growth %'
for c in cols: ws[f'{c}16']=0.08

ws['A17']='Borrowings growth %'
borg = {'D':0.08,'E':0.07,'F':0.06,'G':0.05,'H':0.05,'I':0.05}
for c,v in borg.items(): ws[f'{c}17']=v

ws['A18']='Receivable days (DSO)'
dso = {'D':62,'E':60,'F':58,'G':57,'H':56,'I':55}
for c,v in dso.items(): ws[f'{c}18']=v

ws['A19']='Inventory days'
for c in cols: ws[f'{c}19']=19

ws['A20']='Payable days (DPO, on revenue)'
for c in cols: ws[f'{c}20']=46

ws['A21']='Minority interest growth %'
for c in cols: ws[f'{c}21']=0

# ---------- B. REVENUE BUILD (rows 24-30) ----------
ws['A24']='Organic revenue (prior total x (1+organic growth))'
ws['C24']='=Historicals!K6'
ws['D24']='=C30*(1+D4)'
for i in range(1,6):
    c=cols[i]; p=cols[i-1]
    ws[f'{c}24']=f'={p}30*(1+{c}4)'
ws['A25']='  Organic growth %'
for c in cols: ws[f'{c}25']=f'={c}4'
ws['A26']='Inorganic (acquired) revenue added'
ws['C26']=0
for c in cols: ws[f'{c}26']=f'={c}5'
ws['A27']='  Total revenue growth %'
ws['D27']='=D30/C30-1'
for i in range(1,6):
    c=cols[i]; p=cols[i-1]
    ws[f'{c}27']=f'={c}30/{p}30-1'
# Total Operating Revenues (row30)
ws['C30']='=Historicals!K6'
for c in cols: ws[f'{c}30']=f'={c}24+{c}26'

# ---------- C. INCOME STATEMENT tweaks ----------
# cost of revenue on operating revenue (row30), opex on operating revenue
for c in cols:
    ws[f'{c}34']=f'={c}30*{c}6'      # cost of revenue
    ws[f'{c}37']=f'={c}30*{c}9'      # SG&A
# EBITDA (row51) = operating revenue - cost - SG&A ; margin on operating revenue
for c in ['C']+cols:
    ws[f'{c}51']=f'={c}30-{c}34-{c}37'
    ws[f'{c}52']=f'={c}51/{c}30'
# EPS share count -> DCF!B18 (18.0cr)
for c in ['C']+cols:
    ws[f'{c}49']=f'={c}46/DCF!$B$18'

# relabel CF NBFC row
ws['A93']='Change in other loans/advances'

# IS line relabels
ws['A34']='Cost of Revenue (hardware, materials & subcontracting)'
ws['A37']='Employee & Other Operating Expenses (SG&A)'

wb.save('Black_Box_Co_Model.xlsx')
print('Forecast rebuilt & saved.')
